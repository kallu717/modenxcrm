"""Push a Lead Desk "Export for Pipeline" file into the Pipeline API.

    python scripts/push_leads.py modenx-pipeline-import-20261022.xlsx \
        --url https://pipeline.example.internal --dry-run
    python scripts/push_leads.py leads.csv --url https://pipeline.example.internal

The API key comes from the PIPELINE_INGEST_KEY environment variable (never a
command-line argument, so it stays out of shell history and process lists).
Leads arrive unassigned; a manager assigns accounts to reps in Pipeline.
Duplicates (by email, against every lead already in Pipeline) are skipped, so
re-running the same file is safe.

Reads .csv, or .xlsx (needs `pip install openpyxl`).
"""

import argparse
import csv
import json
import os
import sys
import urllib.error
import urllib.request
from pathlib import Path

COLUMNS = {  # Lead Desk / Pipeline template header -> API field
    "contact name": "name", "job title": "title", "account name": "account",
    "account type": "type", "email": "email", "email verified": "verified",
    "phone": "phone", "city": "city", "notes": "notes",
}
BATCH = 500


def read_rows(path: Path) -> list[dict]:
    if path.suffix.lower() == ".csv":
        with path.open(newline="", encoding="utf-8-sig") as f:
            raw = list(csv.DictReader(f))
    elif path.suffix.lower() in (".xlsx", ".xlsm"):
        try:
            from openpyxl import load_workbook
        except ImportError:
            sys.exit("Reading .xlsx needs openpyxl: pip install openpyxl (or save the sheet as .csv)")
        ws = load_workbook(path, read_only=True, data_only=True).worksheets[0]
        it = ws.iter_rows(values_only=True)
        header = [str(h or "").strip() for h in next(it, [])]
        raw = [dict(zip(header, row)) for row in it if any(c not in (None, "") for c in row)]
    else:
        sys.exit("Give a .csv or .xlsx file")
    rows = []
    for r in raw:
        row = {}
        for k, v in r.items():
            field = COLUMNS.get(str(k).strip().lower())
            if field:
                row[field] = "" if v is None else str(v).strip()
        rows.append(row)
    return rows


def post(url: str, key: str, body: dict) -> dict:
    req = urllib.request.Request(
        url, data=json.dumps(body).encode(), method="POST",
        headers={"Content-Type": "application/json", "Authorization": f"Bearer {key}"},
    )
    try:
        with urllib.request.urlopen(req, timeout=60) as r:
            return json.load(r)
    except urllib.error.HTTPError as e:
        sys.exit(f"Pipeline refused the request: HTTP {e.code} {e.read().decode(errors='replace')[:300]}")
    except urllib.error.URLError as e:
        sys.exit(f"Couldn't reach Pipeline: {e.reason}")


def main() -> None:
    ap = argparse.ArgumentParser(description=__doc__, formatter_class=argparse.RawDescriptionHelpFormatter)
    ap.add_argument("file", type=Path)
    ap.add_argument("--url", required=True, help="Pipeline base URL, e.g. https://pipeline.example.internal")
    ap.add_argument("--dry-run", action="store_true", help="validate and count, change nothing")
    args = ap.parse_args()

    key = os.environ.get("PIPELINE_INGEST_KEY", "")
    if not key:
        sys.exit("Set PIPELINE_INGEST_KEY in the environment")
    rows = read_rows(args.file)
    if not rows:
        sys.exit("No data rows found")

    endpoint = args.url.rstrip("/") + "/ingest/leads" + ("?dry_run=true" if args.dry_run else "")
    totals = {"created": 0, "ready": 0, "newAccounts": 0, "duplicates": 0, "rejected": 0}
    issues: list[str] = []
    for i in range(0, len(rows), BATCH):
        res = post(endpoint, key, {"filename": args.file.name, "source": "Lead Desk", "rows": rows[i:i + BATCH]})
        for k in totals:
            totals[k] += res[k]
        # row numbers in messages are per batch; shift them to the file's row numbers
        issues += [m.replace("Row ", f"Row {i}+", 1) if i else m for m in res["issues"]]

    print(("DRY RUN — nothing was saved\n" if args.dry_run else "") +
          f"{len(rows)} rows read: {totals['ready']} valid, {totals['created']} created, "
          f"{totals['duplicates']} duplicates skipped, {totals['rejected']} rejected, "
          f"{totals['newAccounts']} new accounts")
    for m in issues[:50]:
        print("  " + m)
    if len(issues) > 50:
        print(f"  …and {len(issues) - 50} more")


if __name__ == "__main__":
    main()
